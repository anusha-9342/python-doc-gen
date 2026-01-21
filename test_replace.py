from docx import Document
from placeholder_replacer import replace_placeholders

doc = Document("template.docx")

mapping = {
    "{{company_name}}": "Acme Corp",
    "{{customer_name}}": "Anusha",
    "{{product_id}}": "P-1009",
    "{{date}}": "20-01-2026"
}

replace_placeholders(doc, mapping)
doc.save("output.docx")

print("Header preservation test completed.")
