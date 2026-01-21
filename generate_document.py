import os
import pandas as pd
from docx import Document
from placeholder_replacer import replace_placeholders  # Your function

# Path to your template Word document
TEMPLATE_PATH = "template.docx"

# Path to input spreadsheet
SPREADSHEET_PATH = "datasets/sample.xlsx"  

# Output directory
OUTPUT_DIR = "generated_docs"

# Placeholder mapping: template placeholder -> spreadsheet column
PLACEHOLDER_MAPPING = {
    "{{customer_name}}": "Customer Name",
    "{{product_id}}": "Product ID",
    "{{date}}": "Date"
}


def create_output_dir(path):
    """Create output directory if it doesn't exist."""
    if not os.path.exists(path):
        os.makedirs(path)

def generate_filename(row):
    """
    Construct a unique and descriptive filename for each document
    Example: 'Invoice_Anusha_P-1001.docx'
    """
    customer = str(row["Customer Name"]).replace(" ", "_")
    product = str(row["Product ID"])
    return f"Invoice_{customer}_{product}.docx"

def row_to_mapping(row):
    """
    Convert a pandas DataFrame row to a placeholder->value dict
    according to PLACEHOLDER_MAPPING
    """
    return {placeholder: row[column] for placeholder, column in PLACEHOLDER_MAPPING.items()}


def main():
    #  Load spreadsheet into pandas
    if SPREADSHEET_PATH.endswith(".xlsx"):
        df = pd.read_excel(SPREADSHEET_PATH)
    else:
        df = pd.read_csv(SPREADSHEET_PATH)

    #  Create output directory
    create_output_dir(OUTPUT_DIR)

    #  Loop through each row
    for index, row in df.iterrows():
        #  Load template
        doc = Document(TEMPLATE_PATH)

        # Convert row to placeholder->value mapping
        mapping = row_to_mapping(row)

        # Replace placeholders in body & header
        replace_placeholders(doc, mapping)

        # 5️⃣ Generate unique filename
        filename = generate_filename(row)
        output_path = os.path.join(OUTPUT_DIR, filename)

        # Save the populated document
        doc.save(output_path)

        print(f"[{index+1}/{len(df)}] Generated: {output_path}")

    print("All documents generated successfully!")



if __name__ == "__main__":
    main()
