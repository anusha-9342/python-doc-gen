from docx import Document
from placeholder_replacer import replace_placeholders

# Load the template
doc = Document("template.docx")

# Test data (mock values)
test_mapping = {
    "{{customer_name}}": "Anusha",
    "{{product_id}}": "P-1023",
    "{{date}}": "20-01-2026"
}

# Run replacement
replace_placeholders(doc, test_mapping)

# Save output
doc.save("output.docx")

print("Test completed. Check output.docx")
